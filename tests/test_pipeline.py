import sys
import pandas as pd

import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from hrdash.attendance.parser import clean_attendance_data, clean_approved_leave
from hrdash.employee_master import load_employee_master, merge_master_into_log
from hrdash.attendance import engine as att_engine
from hrdash.attendance import analytics as att_analytics
from hrdash.config import AppConfig, AttendanceConfig, PayrollConfig, config_to_dataframe, config_from_dataframe
from hrdash.payroll import calculate_payroll
from hrdash.sales import clean_sales_activity, calculate_sales_performance, combine_attendance_and_sales
from hrdash.reporting.excel_report import ReportData, generate_workbook

# --- Raw attendance data inspired by the spec's example table ---
# 2026-08-24 = Monday ... 2026-08-29 = Saturday, 2026-08-30 = Sunday.
rows = [
    ("OUR COMPANY", "1111", "1111", "24/08/2026 07.34.23"),
    ("HRD & ADMIN", "", "12021007", "24/08/2026 07.40.16"),
    ("HRD & ADMIN", "", "12021007", "24/08/2026 17.24.11"),
    ("HRD & ADMIN", "", "12021007", "25/08/2026 07.47.28"),
    ("HRD & ADMIN", "", "12021007", "25/08/2026 17.25.02"),
    ("HRD & ADMIN", "", "12021007", "26/08/2026 14.22.01"),  # single scan afternoon (Wed)
    ("HRD & ADMIN", "", "12021007", "27/08/2026 07.58.35"),
    ("HRD & ADMIN", "", "12021007", "27/08/2026 17.03.16"),
    ("HRD & ADMIN", "", "12021007", "28/08/2026 07.54.08"),
    ("HRD & ADMIN", "", "12021007", "28/08/2026 17.01.28"),
    ("HRD & ADMIN", "", "12021007", "28/08/2026 17.24.34"),  # triple scan day (Fri)
    ("HRD & ADMIN", "", "12021007", "29/08/2026 07.51.57"),  # Saturday, working
    ("Sales Mobil Baru", "Budi", "EMP001", "24/08/2026 07.50.00"),
    ("Sales Mobil Baru", "Budi", "EMP001", "24/08/2026 17.10.00"),
    ("Sales Mobil Baru", "Budi", "EMP001", "25/08/2026 08.30.00"),  # late (Tue)
    ("Sales Mobil Baru", "Budi", "EMP001", "25/08/2026 17.10.00"),
    # 26/08 (Wed) no scan at all for EMP001 -> Mangkir
    ("Sales Mobil Baru", "Budi", "EMP001", "30/08/2026 09.00.00"),  # Sunday off -> holiday scan
]
raw = pd.DataFrame(rows, columns=["Department", "Name", "No.", "Date/Time"])

clean, rejected, quality = clean_attendance_data(raw)
print("quality:", quality)
assert quality["valid_rows"] == len(raw)

master_raw = pd.DataFrame(
    {
        "No.": ["EMP001", "12021007"],
        "Name": ["Budi Santoso", "Ani Lestari"],
        "Department": ["Sales Mobil Baru", "HRD & Admin"],
        "Position": ["Sales Executive", "Staff HRD"],
        "Employee_Type": ["SALES", "OFFICE"],
        "Monthly_Salary": [5500000, 6000000],
        "Active": ["Y", "Y"],
    }
)
master = load_employee_master(master_raw)
print("\nmaster:\n", master)

enriched, unmapped = merge_master_into_log(clean, master)
print("\nunmapped:", unmapped.to_dict("records"))

cfg = AppConfig(attendance=AttendanceConfig(saturday_is_working=False))  # so 29/08 Sat is a holiday scan
daily = att_engine.build_daily_attendance(enriched, cfg.attendance, holidays=None, approved_leave=None)
print("\ndaily sample:\n", daily[["No.", "Tanggal", "Hari", "Status_Masuk", "Status_Pulang", "Scan_Count", "Anomaly_Type", "Severity"]].to_string())

# Sanity checks against the spec's worked example (Section 17)
row_26 = daily[(daily["No."] == "12021007") & (daily["Tanggal"].astype(str) == "2026-08-26")].iloc[0]
assert row_26["Scan_Count"] == 1
assert row_26["Status_Masuk"] == "Lupa Absen Masuk"  # single afternoon scan -> missing clock-in
print("\n26/08 single-scan row ->", row_26["Status_Masuk"], row_26["Status_Pulang"], row_26["Anomaly_Type"], row_26["Severity"])

row_28 = daily[(daily["No."] == "12021007") & (daily["Tanggal"].astype(str) == "2026-08-28")].iloc[0]
assert row_28["Scan_Count"] == 3
assert row_28["Jam_Masuk"] == "07:54:08"
assert row_28["Jam_Pulang"] == "17:24:34"
print("28/08 triple-scan row ->", row_28["Jam_Masuk"], row_28["Jam_Pulang"], row_28["Anomaly_Type"])

row_mangkir = daily[(daily["No."] == "EMP001") & (daily["Tanggal"].astype(str) == "2026-08-26")].iloc[0]
print("EMP001 26/08 (no scan) ->", row_mangkir["Status_Masuk"], row_mangkir["Severity"])
assert row_mangkir["Status_Masuk"] == "Mangkir"

row_sun = daily[(daily["No."] == "EMP001") & (daily["Tanggal"].astype(str) == "2026-08-30")].iloc[0]
print("EMP001 30/08 (Sunday off, scanned) ->", row_sun["Status_Masuk"], row_sun["Anomaly_Type"])
assert row_sun["Anomaly_Type"] == "Scan Pada Hari Libur"

employee_summary = att_analytics.calculate_employee_summary(daily, cfg.score)
department_summary = att_analytics.calculate_department_summary(daily)
trend = att_analytics.calculate_daily_trend(daily)
anomalies = daily[daily["Is_Anomali"] == 1].copy()
insights = att_analytics.derive_hr_insights(employee_summary, department_summary)
print("\nemployee_summary:\n", employee_summary.to_string())
print("\ninsights:", insights)

# Payroll: without a configured rate, deductions must stay NA, not 0.
payroll_unconfigured = calculate_payroll(employee_summary, master, PayrollConfig())
assert payroll_unconfigured["Potongan_Telat"].isna().all()
print("\npayroll (unconfigured rates):\n", payroll_unconfigured.to_string())

payroll_configured = calculate_payroll(
    employee_summary, master, PayrollConfig(late_deduction_per_minute=5000, absence_deduction_per_day=250000)
)
print("\npayroll (configured rates):\n", payroll_configured.to_string())
assert not payroll_configured["Potongan_Telat"].isna().all()

# Test late deduction per occurrence & sales missing clock-out penalty
payroll_occ = calculate_payroll(
    employee_summary, master, PayrollConfig(
        late_deduction_mode="per_occurrence",
        late_deduction_per_occurrence=50000,
        sales_no_clock_out_deduction=100000,
    )
)
print("\npayroll (per occurrence & sales penalty):\n", payroll_occ.to_string())
assert "Jumlah_Tidak_Absen_Pulang" in payroll_occ.columns
assert "Potongan_Lupa_Pulang_Sales" in payroll_occ.columns
# EMP001 is SALES; non-sales should have 0 or NA for Potongan_Lupa_Pulang_Sales
sales_row_p = payroll_occ[payroll_occ["No."] == "EMP001"].iloc[0]
office_row_p = payroll_occ[payroll_occ["No."] == "12021007"].iloc[0]
assert office_row_p["Potongan_Lupa_Pulang_Sales"] == 0.0

# Multi-sheet leave template check
from hrdash.leave import build_leave_template, load_leave_workbook
leave_tpl_bytes = build_leave_template()
assert len(leave_tpl_bytes) > 0

# Sales activity / performance (independent of fingerprint scans)
sales_raw = pd.DataFrame(
    {
        "No.": ["EMP001", "EMP001"],
        "Tanggal": ["21/08/2026", "22/08/2026"],
        "Prospect": [8, 6],
        "Visit": [4, 3],
        "Test Drive": [2, 1],
        "SPK": [1, 1],
        "Delivery": [0, 1],
        "Revenue": [0, 250000000],
    }
)
sales_activity = clean_sales_activity(sales_raw)
sales_perf = calculate_sales_performance(sales_activity, master, cfg.sales_score)
print("\nsales_performance:\n", sales_perf.to_string())
combined = combine_attendance_and_sales(employee_summary, sales_perf)
print("\ncombined attendance+sales:\n", combined.to_string())

# Config round-trip
cfg_df = config_to_dataframe(cfg)
cfg2 = config_from_dataframe(cfg_df)
assert cfg2.attendance.grace_minutes == cfg.attendance.grace_minutes
print("\nconfig round-trip OK")

# Full workbook generation
report = ReportData(
    config=cfg,
    employee_master=master,
    raw_log=clean,
    daily_attendance=daily,
    employee_summary=employee_summary,
    department_summary=department_summary,
    daily_trend=trend,
    anomalies=anomalies,
    rejected=rejected,
    payroll=payroll_configured,
    sales_activity=sales_activity,
    sales_performance=sales_perf,
    insights=insights,
    report_period_label="2026-08-21 s/d 2026-08-29",
)
wb_bytes = generate_workbook(report)
with open("/tmp/Laporan_HR_Test.xlsx", "wb") as f:
    f.write(wb_bytes)
print("\nWorkbook written:", len(wb_bytes), "bytes")

from openpyxl import load_workbook
wb = load_workbook("/tmp/Laporan_HR_Test.xlsx")
print("Sheets:", wb.sheetnames)
assert "DASHBOARD" in wb.sheetnames
assert "CONFIG" in wb.sheetnames
assert "SALES_PERFORMANCE" in wb.sheetnames
print("\nALL SMOKE TESTS PASSED")
