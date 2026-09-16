"""
Employee Master (Section 9).

Employee Type detection: if ``Department`` contains "SALES" the employee
defaults to type SALES, otherwise OFFICE — but HR can always override this
via the ``Employee_Type`` column in the uploaded master file, as required
by the spec.
"""

from __future__ import annotations

import io
import datetime as dt
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

EMPLOYEE_MASTER_COLUMNS = [
    "No.",
    "Name",
    "Department",
    "Position",
    "Employee_Type",
    "Jam_Masuk",
    "Jam_Pulang",
    "Monthly_Salary",
    "Active",
]

VALID_EMPLOYEE_TYPES = {"OFFICE", "SALES"}

NAVY = "1F4E78"
WHITE = "FFFFFF"


class EmployeeMasterError(ValueError):
    """Raised when the uploaded employee master file is malformed."""


def _clean_col(name: object) -> str:
    return str(name).replace("\n", " ").strip()


def _parse_time_str(val: object) -> Optional[str]:
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (dt.time,)):
        return val.strftime("%H:%M")
    if isinstance(val, (pd.Timestamp, dt.datetime)):
        return val.strftime("%H:%M")
    s = str(val).strip()
    if not s or s.lower() in ("nan", "none", "nat", ""):
        return None
    s = s.replace(".", ":")
    parts = s.split(":")
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return None
    return None


def detect_employee_type(department: Optional[str]) -> str:
    if department and "sales" in str(department).lower():
        return "SALES"
    return "OFFICE"


def load_employee_master(raw_df: Optional[pd.DataFrame]) -> pd.DataFrame:
    """Validate and normalize an uploaded Employee Master file.

    Returns a DataFrame with exactly ``EMPLOYEE_MASTER_COLUMNS``, one row
    per unique Employee ID (last occurrence wins so HR can re-upload an
    updated file). Missing optional fields are filled with safe defaults;
    ``No.`` is mandatory.
    """
    if raw_df is None or raw_df.empty:
        return pd.DataFrame(columns=EMPLOYEE_MASTER_COLUMNS)

    df = raw_df.copy()
    df.columns = [_clean_col(c) for c in df.columns]

    id_col = next((c for c in ["No.", "Employee ID", "ID", "NIK"] if c in df.columns), None)
    if id_col is None:
        raise EmployeeMasterError(
            "Employee Master harus memiliki kolom employee ID (No./Employee ID/ID/NIK)."
        )

    out = pd.DataFrame()
    out["No."] = df[id_col].astype("string").str.strip()
    out["Name"] = df["Name"].astype("string").str.strip() if "Name" in df.columns else pd.NA
    out["Department"] = (
        df["Department"].astype("string").str.strip() if "Department" in df.columns else pd.NA
    )
    out["Position"] = df["Position"].astype("string").str.strip() if "Position" in df.columns else pd.NA

    if "Employee_Type" in df.columns:
        etype = df["Employee_Type"].astype("string").str.strip().str.upper()
    elif "Employee Type" in df.columns:
        etype = df["Employee Type"].astype("string").str.strip().str.upper()
    else:
        etype = pd.Series([pd.NA] * len(df), dtype="string")
    out["Employee_Type"] = etype

    # Working hours per employee / division
    in_col = next(
        (c for c in ["Jam_Masuk", "Jam Masuk", "Jam_Kerja_Masuk", "Jam Kerja Masuk", "In", "Clock_In", "Clock In"] if c in df.columns),
        None,
    )
    out_col = next(
        (c for c in ["Jam_Pulang", "Jam Pulang", "Jam_Kerja_Pulang", "Jam Kerja Pulang", "Out", "Clock_Out", "Clock Out"] if c in df.columns),
        None,
    )
    range_col = next(
        (c for c in ["Jam_Kerja", "Jam Kerja", "Range_Jam_Kerja", "Schedule", "Jadwal"] if c in df.columns),
        None,
    )

    if in_col is not None:
        out["Jam_Masuk"] = df[in_col].map(_parse_time_str).astype("string")
    elif range_col is not None:
        def _get_start(val):
            if val is None or pd.isna(val):
                return None
            txt = str(val)
            for sep in ["-", "s/d", "to"]:
                if sep in txt:
                    return _parse_time_str(txt.split(sep)[0])
            return _parse_time_str(txt)
        out["Jam_Masuk"] = df[range_col].map(_get_start).astype("string")
    else:
        out["Jam_Masuk"] = pd.Series([pd.NA] * len(df), dtype="string")

    if out_col is not None:
        out["Jam_Pulang"] = df[out_col].map(_parse_time_str).astype("string")
    elif range_col is not None:
        def _get_end(val):
            if val is None or pd.isna(val):
                return None
            txt = str(val)
            for sep in ["-", "s/d", "to"]:
                if sep in txt:
                    return _parse_time_str(txt.split(sep)[1])
            return None
        out["Jam_Pulang"] = df[range_col].map(_get_end).astype("string")
    else:
        out["Jam_Pulang"] = pd.Series([pd.NA] * len(df), dtype="string")

    if "Monthly_Salary" in df.columns:
        salary_col = df["Monthly_Salary"]
    elif "Monthly Salary" in df.columns:
        salary_col = df["Monthly Salary"]
    else:
        salary_col = pd.Series([pd.NA] * len(df))
    out["Monthly_Salary"] = pd.to_numeric(salary_col, errors="coerce")

    if "Active" in df.columns:
        active_raw = df["Active"].astype("string").str.strip().str.lower()
        out["Active"] = ~active_raw.isin(["n", "no", "tidak", "0", "false", "inactive"])
    else:
        out["Active"] = True

    out = out.dropna(subset=["No."])
    out = out[out["No."] != ""]

    # Fill employee type using department heuristic where HR left it blank.
    needs_default = out["Employee_Type"].isna() | ~out["Employee_Type"].isin(VALID_EMPLOYEE_TYPES)
    out.loc[needs_default, "Employee_Type"] = out.loc[needs_default, "Department"].map(
        detect_employee_type
    )

    out["Name"] = out["Name"].fillna("").replace("", pd.NA)
    out["Name"] = out["Name"].fillna(out["No."].map(lambda x: f"Karyawan {x}"))
    out["Department"] = out["Department"].fillna("Belum Dipetakan")
    out["Position"] = out["Position"].fillna("")

    out = out.drop_duplicates(subset=["No."], keep="last").reset_index(drop=True)
    return out[EMPLOYEE_MASTER_COLUMNS]


def merge_master_into_log(
    clean_log: pd.DataFrame, master: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Enrich attendance rows with employee master identity/type/schedule.

    Where an employee exists in the raw log but not in the master file, the
    identity fields already present on the log (or safe fallbacks) are kept
    and the employee is flagged in the returned "unmapped" DataFrame so HR
    can extend the master file. No attendance rows are dropped.
    """
    if master is None or master.empty:
        enriched = clean_log.copy()
        enriched["Employee_Type"] = enriched["Department"].map(detect_employee_type)
        enriched["Position"] = ""
        enriched["Jam_Masuk_Master"] = pd.NA
        enriched["Jam_Pulang_Master"] = pd.NA
        unmapped = enriched[["No.", "Name", "Department"]].drop_duplicates()
        return enriched, unmapped

    m = master.set_index("No.")
    enriched = clean_log.copy()
    in_master = enriched["No."].isin(m.index)

    enriched["Employee_Type"] = enriched["No."].map(m["Employee_Type"]).astype("string")
    enriched["Position"] = enriched["No."].map(m["Position"]).astype("string").fillna("")
    if "Jam_Masuk" in m.columns:
        enriched["Jam_Masuk_Master"] = enriched["No."].map(m["Jam_Masuk"]).astype("string")
    else:
        enriched["Jam_Masuk_Master"] = pd.NA
    if "Jam_Pulang" in m.columns:
        enriched["Jam_Pulang_Master"] = enriched["No."].map(m["Jam_Pulang"]).astype("string")
    else:
        enriched["Jam_Pulang_Master"] = pd.NA

    # Master identity wins for name/department when available (single source
    # of truth), otherwise fall back to whatever the machine log carried.
    master_name = enriched["No."].map(m["Name"])
    master_dept = enriched["No."].map(m["Department"])
    enriched["Name"] = master_name.fillna(enriched["Name"])
    enriched["Department"] = master_dept.fillna(enriched["Department"])
    enriched["Employee_Type"] = enriched["Employee_Type"].fillna(
        enriched["Department"].map(detect_employee_type)
    )

    unmapped = (
        enriched.loc[~in_master, ["No.", "Name", "Department"]]
        .drop_duplicates()
        .reset_index(drop=True)
    )
    return enriched, unmapped


def build_employee_master_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Master"
    guide = wb.create_sheet("Petunjuk")

    ws.append(EMPLOYEE_MASTER_COLUMNS)
    ws.append(["EMP001", "Budi Santoso", "Sales Mobil Baru", "Sales Executive", "SALES", "08:30", "17:30", 5500000, "Y"])
    ws.append(["EMP002", "Sari Wijaya", "HRD & Admin", "Staff HRD", "OFFICE", "08:00", "17:00", 6000000, "Y"])

    header_fill = PatternFill("solid", fgColor=NAVY)
    header_font = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    body_font = Font(name="Segoe UI", size=10, color="1F2937")
    thin = Side(style="thin", color="D1D5DB")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    for row in ws.iter_rows(min_row=2, max_row=1000, max_col=len(EMPLOYEE_MASTER_COLUMNS)):
        for cell in row:
            cell.font = body_font
            cell.border = Border(bottom=thin)

    dv_type = DataValidation(type="list", formula1='"OFFICE,SALES"', allow_blank=True)
    dv_type.error = "Pilih OFFICE atau SALES."
    dv_type.errorTitle = "Employee Type tidak valid"
    ws.add_data_validation(dv_type)
    dv_type.add("E2:E1000")

    dv_active = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv_active)
    dv_active.add("I2:I1000")

    widths = {"A": 14, "B": 22, "C": 20, "D": 20, "E": 14, "F": 14, "G": 14, "H": 16, "I": 10}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I1000"

    guide_rows = [
        ["PETUNJUK EMPLOYEE MASTER"],
        ["No.", "Employee ID — harus sama dengan kolom No. pada export mesin absensi."],
        ["Name", "Nama karyawan."],
        ["Department", "Departemen/divisi."],
        ["Position", "Jabatan (opsional)."],
        ["Employee_Type", "OFFICE atau SALES. Jika dikosongkan, sistem menebak dari Department (mengandung 'SALES')."],
        ["Jam_Masuk", "Jam masuk resmi per divisi/karyawan (format HH:MM, contoh: 08:30). Jika dikosongkan, mengikuti jadwal dashboard."],
        ["Jam_Pulang", "Jam pulang resmi per divisi/karyawan (format HH:MM, contoh: 17:30). Jika dikosongkan, mengikuti jadwal dashboard."],
        ["Monthly_Salary", "Gaji bulanan untuk perhitungan payroll (opsional, angka saja)."],
        ["Active", "Y = aktif, N = non-aktif. Karyawan non-aktif tetap dihitung di riwayat tapi ditandai."],
    ]
    for row in guide_rows:
        guide.append(row)
    guide["A1"].font = Font(name="Segoe UI", size=14, bold=True, color=WHITE)
    guide["A1"].fill = header_fill
    guide.merge_cells("A1:B1")
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 90

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
